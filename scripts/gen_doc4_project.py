"""
Generate Document 4: Project Documentation for EMS Pro.
Covers Project Scope, Business Requirements, Work Packages (WP-01 to WP-29),
Seeded Demo Data, Feature Inventories, and Deployment Handover.
Outputs to: /Users/apple/Downloads/EMS/document/4_Project_Documentation.docx
"""
import os
import sys
from docx.shared import Inches, Pt, RGBColor
from scripts.docx_helper import (
    create_styled_document, add_cover_page, add_header_1, add_header_2,
    add_header_3, add_callout, add_styled_table, COLOR_DARK_TEXT
)

def build_project_documentation():
    doc = create_styled_document()
    add_cover_page(
        doc,
        title="EMS Pro — Project Documentation",
        subtitle="Business Scope, Work Package Completion Audit, Functional Module Inventory & Handover Guide",
        doc_type="Product, Project & Operational Handover Manual"
    )

    # 1. Project Background & Vision
    add_header_1(doc, "1. Project Vision & Executive Summary")
    doc.add_paragraph(
        "EMS Pro is a modern, enterprise-scale Human Resource Management, Statutory Payroll, "
        "Performance Management, and Projects SaaS platform. "
        "The project was conceived to solve the common challenges of fragmented corporate operations: "
        "disparate spreadsheets for employee records, error-prone manual payroll calculations, disconnected attendance devices, "
        "opaque annual review processes, and lack of real-time project profitability and timesheet tracking."
    )
    doc.add_paragraph(
        "By consolidating these mission-critical business workflows into a unified, secure, multi-tenant architecture, "
        "EMS Pro empowers enterprise HR admins, managers, and employees to interact through intuitive role-tailored dashboards "
        "while guaranteeing zero cross-company data leakage."
    )

    # 2. Key Business Objectives & Scope
    add_header_1(doc, "2. Key Business Objectives")
    objectives = [
        ("Multi-Tenant Cloud Scalability", "Support dozens to thousands of autonomous corporate organizations on a unified infrastructure with complete data isolation."),
        ("Statutory Compliance Automation", "Automate compliance with Indian labor regulations (EPF Act, ESI Act, State Professional Tax rules, and Income Tax TDS slabs) with zero mathematical discrepancies."),
        ("Comprehensive Employee Lifecycle", "Manage seamless transitions from pre-onboarding, employment verification, shift attendance, regular leave cycles, promotions, to resignation and Full & Final (FnF) settlements."),
        ("Integrated Work & Project Accounting", "Bridge employee operational time logs directly to project billable budgets, task kanbans, and manager approval queues."),
        ("Auditable Enterprise Transparency", "Provide transparent self-service payslips, tax sheets, leave balances, and company announcements with tamper-proof audit trails."),
    ]
    for obj_title, obj_desc in objectives:
        p = doc.add_paragraph()
        r = p.add_run(f"• {obj_title}: ")
        r.font.bold = True
        p.add_run(obj_desc)

    # 3. Work Package Delivery Audit (WP-01 to WP-29)
    add_header_1(doc, "3. Complete Work Package Delivery Matrix")
    doc.add_paragraph(
        "The platform development was planned and executed across 29 discrete Work Packages (WP-01 to WP-29), "
        "all of which have been implemented, tested, and verified in production:"
    )

    wp_headers = ["WP Code", "Module / Area", "Delivered Capabilities & Key Features", "Status"]
    wp_data = [
        ["WP-01", "Core Architecture", "PostgreSQL RLS infrastructure, Alembic migrations, base schemas, Celery setup", "Delivered & Verified"],
        ["WP-02", "Identity & Auth", "Argon2id password hashing, JWT stateless access tokens, refresh token rotation", "Delivered & Verified"],
        ["WP-03", "Company Onboarding", "Self-registration, Super Admin approval queue, domain assignment, company settings", "Delivered & Verified"],
        ["WP-04", "Tenant Isolation Audit", "Cross-tenant boundary test suite, RLS kernel validation, leakage prevention", "Delivered & Verified"],
        ["WP-05", "Core HR - Employees", "Master employee records, employee codes (INF-001), hierarchy, department mapping", "Delivered & Verified"],
        ["WP-06", "Departments & Org", "Department creation, manager assignments, reporting hierarchy navigation", "Delivered & Verified"],
        ["WP-07", "Attendance Engine", "Web check-in/out, biometric sync agent, geofencing, daily hours auto-calculation", "Delivered & Verified"],
        ["WP-08", "Shift Management", "Rotational shifts, flexible work schedules, night shift allowance computation", "Delivered & Verified"],
        ["WP-09", "Holiday Calendars", "Company & department-specific holiday lists, mandatory vs optional holidays", "Delivered & Verified"],
        ["WP-10", "Leave Types & Balances", "Annual quota allocation, carry-forward rules, leave balance ledger", "Delivered & Verified"],
        ["WP-11", "Leave Requests & Approvals", "Multi-stage leave approval state machine, email notifications, manager alerts", "Delivered & Verified"],
        ["WP-12", "HR Analytics Dashboard", "Real-time attendance percentages, active headcount, leave trends, pending tasks", "Delivered & Verified"],
        ["WP-13", "Employee Self-Service", "Personal profile editing, emergency contacts, leave balances, timesheets", "Delivered & Verified"],
        ["WP-14", "Manager Team Workspace", "Team attendance roster, shift assignments, pending leave approvals", "Delivered & Verified"],
        ["WP-15", "Super Admin Platform", "Global company directory, platform analytics, system-wide settings", "Delivered & Verified"],
        ["WP-16", "Salary Structures", "Basic, HRA, Special Allowances, percentage/fixed formula component definitions", "Delivered & Verified"],
        ["WP-17", "Statutory Configuration", "EPF 12% caps, ESI threshold rules, State Professional Tax slabs", "Delivered & Verified"],
        ["WP-18", "Employee Salary Assignment", "CTC breakdown, effective-from dates, increment tracking, revision history", "Delivered & Verified"],
        ["WP-19", "Automated Payroll Runs", "Monthly batch payroll execution, attendance-based salary proration, deductions", "Delivered & Verified"],
        ["WP-20", "Payslip Generation", "Detailed payslips, earnings/deductions breakdown, PDF export, self-service download", "Delivered & Verified"],
        ["WP-21", "Reimbursements Engine", "Expense claim submissions, receipt attachment, manager review & payroll payout", "Delivered & Verified"],
        ["WP-22", "Performance Goals", "SMART goal tracking, target metrics, employee progress updates, weightage", "Delivered & Verified"],
        ["WP-23", "Performance Reviews", "Review cycles, self-evaluations, manager reviews, 1-5 rating distribution", "Delivered & Verified"],
        ["WP-24", "Projects & Task Boards", "Project portfolios, interactive Kanban boards, task assignments, due dates", "Delivered & Verified"],
        ["WP-25", "Timesheets & Logged Hours", "Task-level time logging, billable vs non-billable hours, manager approval queue", "Delivered & Verified"],
        ["WP-26", "Platform Services", "Announcements broadcast, HMAC signed file URLs, Document vault, Global Search", "Delivered & Verified"],
        ["WP-27", "Offboarding & FnF", "Resignation submission, notice period calculation, Full & Final financial settlement", "Delivered & Verified"],
        ["WP-28", "Security Hardening", "Rate-limiting buffer, XSS/CORS hardening, RLS leakage regression tests", "Delivered & Verified"],
        ["WP-29", "Final Verification", "Production bundle compilation, complete end-to-end integration tests pass", "Delivered & Verified"],
    ]
    add_styled_table(doc, wp_headers, wp_data, [Inches(0.9), Inches(1.5), Inches(3.1), Inches(1.0)])

    # 4. Multi-Company Production Dataset
    add_header_1(doc, "4. Multi-Tenant Reference Dataset & Seeded Companies")
    doc.add_paragraph(
        "EMS Pro comes pre-seeded with 6 diverse enterprise companies representing different commercial industries, "
        "complete with unique projects, tasks, employees, and payroll configurations:"
    )

    company_headers = ["Company Name", "Industry", "HQ City", "Key Projects", "Workforce Scope"]
    company_data = [
        ["Infiria Systems", "Technology / SaaS", "Ahmedabad", "EMS Enterprise 2.0, Infiria Mobile Suite, AI HR Copilot", "Engineering, Product, Design, HR"],
        ["NexusPay Fintech", "Financial Services", "Mumbai", "Instant UPI Payouts, Automated KYC & AML, Fraud ML", "Risk & Fraud, Engineering, Treasury"],
        ["Aether Cloud Labs", "Cloud & DevOps", "Bengaluru", "K8s Auto-scaler, Edge Gateway, Zero-Trust Perimeter", "Principal SREs, Cloud Architects, Security"],
        ["Zenith Health Tech", "Healthcare", "Hyderabad", "Telehealth Video Consult, EHR HL7 FHIR, Remote Vitals", "Clinical Ops, Tech Leads, QA Leads"],
        ["Solaria Energy Corp", "Manufacturing", "Vadodara", "Solar Inverter, SCADA Grid Telemetry, Yield Predictor", "Plant Ops, Solar R&D, Procurement"],
        ["Bluepeak Demo Tech", "Enterprise Software", "Pune / Global", "Customer Portal, ERP Supply Chain, Enterprise SSO", "25+ demo accounts across 7 depts"],
    ]
    add_styled_table(doc, company_headers, company_data, [Inches(1.5), Inches(1.2), Inches(0.9), Inches(1.8), Inches(1.1)])

    # 5. User Roles & Access Control
    add_header_1(doc, "5. User Roles & Permission Hierarchy")
    roles_headers = ["Role Identifier", "Target User", "Allowed System Capabilities"]
    roles_data = [
        ["super_admin", "Platform Owner (e.g. Neha Vaghela)", "Full platform control. Approve/reject new companies, view cross-company metrics, manage global projects, inspect audit trail."],
        ["hr_admin", "Company HR Director / Admin", "Tenant-level owner. Manage employees, departments, shifts, leave balances, run monthly payroll, configure salary components, approve reimbursements."],
        ["manager", "Team Lead / Engineering Manager", "View and manage direct reports, approve leaves, assign project tasks, approve timesheets, submit performance evaluations."],
        ["employee", "Individual Contributor", "Self-service dashboard, web check-in/out, apply for leave, submit reimbursements, view payslips, update personal goals, log task timesheets."],
    ]
    add_styled_table(doc, roles_headers, roles_data, [Inches(1.3), Inches(1.8), Inches(3.4)])

    # 6. Operational Handover & Production Runbook
    add_header_1(doc, "6. Operations & Maintenance Runbook")
    doc.add_paragraph(
        "Day-to-day operations and administrative maintenance follow standard procedures:"
    )

    runbook_items = [
        ("Daily Automated Attendance Cutoff", "Celery Beat triggers a daily job at 23:59 to mark un-checked-in employees as Absent and flag unclosed shifts."),
        ("Monthly Payroll Processing Window", "HR Admins execute payroll between the 25th and 30th of each month. The run locks working days, settles approved expense claims, and publishes payslips."),
        ("Database Backup & WAL Archiving", "Automated daily pg_dump backups stored offsite with 30-day retention. PostgreSQL WAL archiving ensures continuous point-in-time recovery."),
        ("Secret Rotation", "JWT secret keys, database credentials, and Redis passwords should be rotated bi-annually via the .env configuration without downtime."),
    ]
    for r_title, r_desc in runbook_items:
        add_header_2(doc, r_title)
        doc.add_paragraph(r_desc)

    output_path = "/Users/apple/Downloads/EMS/document/4_Project_Documentation.docx"
    doc.save(output_path)
    print(f"✔ Successfully generated: {output_path}")

if __name__ == "__main__":
    build_project_documentation()
