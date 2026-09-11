"""
Generate Document 2: Architecture Documentation for EMS Pro.
Outputs to: /Users/apple/Downloads/EMS/document/2_Architecture_Documentation.docx
"""
import os
import sys
from docx.shared import Inches, Pt, RGBColor
from scripts.docx_helper import (
    create_styled_document, add_cover_page, add_header_1, add_header_2,
    add_header_3, add_callout, add_styled_table, COLOR_DARK_TEXT
)

def build_architecture_documentation():
    doc = create_styled_document()
    add_cover_page(
        doc,
        title="EMS Pro — Architecture Documentation",
        subtitle="System Topography, Multi-Tenancy Engine, Domain Boundaries & Security Framework",
        doc_type="System & Enterprise Architecture Guide"
    )

    # 1. Architectural Principles & Top-Level Design
    add_header_1(doc, "1. Enterprise Architectural Principles")
    doc.add_paragraph(
        "The architecture of EMS Pro is engineered around strict domain boundaries, zero-trust tenant isolation, "
        "high availability, auditability, and deterministic financial math. "
        "It supports high-concurrency enterprise workloads while providing complete isolation between distinct client companies."
    )

    core_principles = [
        ("Zero-Trust Multi-Tenancy", "Every tenant query is constrained by PostgreSQL Row-Level Security at the database kernel level, preventing logical bugs in application code from crossing company boundaries."),
        ("Stateless Micro-Modular Services", "Business capabilities are isolated into 7 clean modules sharing standard interfaces, allowing independent evolution and straightforward transition to distributed microservices if required."),
        ("Strict Deterministic Financial Math", "No floating-point approximations in payroll or CTC calculations. Python Decimal and PostgreSQL Numeric(12,2) with round_half_up math ensure statutory compliance."),
        ("Immutable Auditability", "Critical actions (payroll runs, salary changes, resignations, company approvals) write append-only audit log records with actor IDs, client IPs, and timestamp hashes."),
        ("Asynchronous Decoupling", "Resource-intensive operations like batch payroll calculation, payslip PDF rendering, and bulk notifications execute asynchronously via Redis and Celery workers without blocking HTTP request threads."),
    ]
    for p_name, p_desc in core_principles:
        p = doc.add_paragraph()
        r_name = p.add_run(f"• {p_name}: ")
        r_name.font.bold = True
        p.add_run(p_desc)

    # 2. System Topography & Physical Architecture
    add_header_1(doc, "2. System Topography & Deployment Architecture")
    doc.add_paragraph(
        "EMS Pro operates as a modern containerized topology deployed across secure VPC tiers:"
    )

    topography_ascii = (
        "                    ┌──────────────────────────────────────────────┐\n"
        "                    │   Client Browser / Mobile PWA (React 18 SPA) │\n"
        "                    └──────────────────────┬───────────────────────┘\n"
        "                                           │ HTTPS / WSS\n"
        "                                           ▼\n"
        "                    ┌──────────────────────────────────────────────┐\n"
        "                    │   Reverse Proxy / API Gateway (Nginx / Cloud) │\n"
        "                    │   SSL Termination & Rate-Limiting Buffer    │\n"
        "                    └──────────────────────┬───────────────────────┘\n"
        "                                           │ ASGI Proxy\n"
        "                                           ▼\n"
        "                    ┌──────────────────────────────────────────────┐\n"
        "                    │      FastAPI Application Cluster (Uvicorn)   │\n"
        "                    │  (Authentication, RBAC, Services, Routers)   │\n"
        "                    └───────┬───────────────────────────────┬──────┘\n"
        "                            │                               │ Task Dispatches\n"
        "                            │ RLS-bound connection pool     ▼\n"
        "                            ▼                     ┌──────────────────┐\n"
        "     ┌──────────────────────────────────────────┐ │ Redis 7 In-Memory│\n"
        "     │       PostgreSQL 16 Engine with RLS      │ │ Cache / Broker   │\n"
        "     │  - 38 Core Schemas & Tables             │ └────────┬─────────┘\n"
        "     │  - Row-Level Security Policies (Tenant)  │          │ Queue Consumptions\n"
        "     │  - Read Replicas & Automated Snapshots   │          ▼\n"
        "     └──────────────────────────────────────────┘ ┌──────────────────┐\n"
        "                                                  │ Celery Workers   │\n"
        "                                                  │ (Payroll/Exports)│\n"
        "                                                  └──────────────────┘\n"
    )
    p_topo = doc.add_paragraph()
    r_topo = p_topo.add_run(topography_ascii)
    r_topo.font.name = "Courier New"
    r_topo.font.size = Pt(7.5)

    # 3. Multi-Tenancy Architecture
    add_header_1(doc, "3. Multi-Tenancy & Row-Level Security Engine")
    doc.add_paragraph(
        "Unlike naive SaaS architectures that filter queries solely by adding 'WHERE company_id = ...' in application code, "
        "EMS Pro adopts a defense-in-depth database RLS architecture. This provides two independent layers of protection:"
    )

    add_header_2(doc, "3.1 Tenant Context Lifecycle")
    doc.add_paragraph(
        "1. Request Authentication: The client provides an Authorization Bearer token. "
        "The JWT contains the verified 'company_id' and 'role' signed by the platform private key.\n"
        "2. Database Session Binding: The FastAPI dependency get_tenant_db executes transaction-local session parameterization:\n"
        "     SELECT set_config('app.current_company_id', '<company-uuid>', true);\n"
        "     SELECT set_config('app.is_platform_admin', 'off', true);\n"
        "3. Kernel Policy Evaluation: Any SELECT, INSERT, UPDATE, or DELETE executed on tenant tables passes through "
        "PostgreSQL's policy evaluation engine. If a query attempts to read or mutate a row whose company_id does not match "
        "app.current_company_id, PostgreSQL automatically omits the row or raises an integrity error.\n"
        "4. Automatic Re-binding on Commit: Because 'is_local=true' clears configuration on transaction commit, "
        "SQLAlchemy's Session 'after_begin' listener intercepts subsequent transactions on the same session and reapplies the tenant context."
    )

    add_header_2(doc, "3.2 Super Admin Platform Isolation Bypass")
    doc.add_paragraph(
        "Platform Super Admins (role: super_admin) have global management privileges across all companies. "
        "When a Super Admin logs in, app.is_platform_admin is set to 'on', allowing global viewing of company metrics, "
        "pending company approvals, and tenant project management without causing tenant isolation conflicts."
    )

    # 4. Domain Subsystem Architecture
    add_header_1(doc, "4. Modular Domain Architecture (7 Core Subsystems)")
    doc.add_paragraph(
        "The platform functionality is decomposed into 7 primary domains, each encapsulated with zero circular dependencies:"
    )

    subsystems = [
        ("Identity & Access Management", "Manages Companies (registration, approval, onboarding), Users (Argon2id auth, 2FA, password resets), Refresh Tokens (family rotation), and RBAC hierarchy (super_admin, hr_admin, manager, employee)."),
        ("Core HR & Organization", "Maintains Employees (master profile, codes, levels L1-L6, hierarchy), Departments (hierarchical structure, department heads), Resignations, Offboarding workflows, and Full & Final (FnF) settlements."),
        ("Time, Attendance & Leave", "Biometric web check-in/out, geofenced logs, Shift scheduling (rotational, night allowance calculation), Holiday calendars, Annual Leave balances, and Multi-tier Leave approval state machines."),
        ("Statutory Payroll Engine", "Configures multi-tier Salary Structures (Basic, HRA, Allowances), Statutory rules (EPF 12% employee/employer, ESI 0.75%/3.25%, Professional Tax state slabs, Tax TDS), automated monthly Payroll Runs, and Reimbursements."),
        ("Performance Management", "Quarterly & Annual Review Cycles, SMART Goal tracking (metrics, targets, progress), 360-degree Reviews (self, manager, HR), Reviewer scoring, and Bell curve performance distribution analysis."),
        ("Projects & Workspace", "Project portfolio management, Team participation, Kanban Task boards (To Do, In Progress, Review, Done), Time Entry timesheets, billable hours tracking, and Project Milestone progression."),
        ("Platform Services", "Broadcast Announcements, Secure File Storage (HMAC SHA-256 signed temporary download links), Encrypted Employee Document vaults, Audit Logs, and Global Multi-Entity Search."),
    ]
    for s_name, s_desc in subsystems:
        add_header_2(doc, s_name)
        doc.add_paragraph(s_desc)

    # 5. Data Flow Architecture
    add_header_1(doc, "5. Core Data Flows & Sequences")
    doc.add_paragraph(
        "Below are architectural sequences for the system's most critical operations:"
    )

    add_header_2(doc, "5.1 Monthly Payroll Calculation & Settlement Flow")
    flow_steps = [
        "1. HR Admin initiates monthly payroll execution via POST /api/v1/payroll/runs.",
        "2. PayrollService locks active employees in the tenant and calculates working days from Attendance and Leave records.",
        "3. Salary engine resolves gross earnings from active EmployeeSalary assignments.",
        "4. Deductions computed: EPF (employee 12% capped), ESI (if gross <= 21,000), Professional Tax (state slab lookup), and TDS.",
        "5. Employer contributions computed: EPF (employer 12% split to EPS 8.33% / EPF 3.67%) and ESI (3.25%).",
        "6. Approved Reimbursements within the pay period are attached and settled.",
        "7. PayrollItem records persisted with immutable JSON snapshots (earnings_json, deductions_json).",
        "8. Run marked as paid; employee payslips unlocked for self-service download.",
    ]
    for st in flow_steps:
        p = doc.add_paragraph(st)
        p.paragraph_format.left_indent = Inches(0.2)

    # 6. Security & Compliance Architecture
    add_header_1(doc, "6. Security Architecture & Compliance Controls")
    sec_headers = ["Security Domain", "Threat Mitigated", "Architectural Control Implemented"]
    sec_data = [
        ["Authentication", "Brute-force & Credential Stuffing", "Argon2id memory-hard hashing, 15-min JWT lifespan, 5-attempt login lockout"],
        ["Token Hijacking", "XSS Access Token Theft", "Refresh token stored exclusively in HttpOnly, SameSite=Lax cookie with rotation"],
        ["Cross-Tenant Leakage", "Tenant data exposure", "PostgreSQL FORCE ROW LEVEL SECURITY with company_id matching policies"],
        ["IDOR / Direct Object Reference", "Unauthorized resource access", "Repository-level company_id ownership check alongside database RLS"],
        ["File Injection / Snooping", "Malicious file execution", "Uploads validated for MIME type, stored in isolated directory with HMAC signed URLs"],
        ["Audit Tampering", "Repudiation of sensitive actions", "Append-only audit_logs table; delete/update triggers blocked"],
        ["Data at Rest", "Physical media compromise", "PostgreSQL database volume encryption, AES-256 for sensitive KYC fields"],
    ]
    add_styled_table(doc, sec_headers, sec_data, [Inches(1.5), Inches(2.2), Inches(2.8)])

    # 7. Scalability & Disaster Recovery
    add_header_1(doc, "7. Scalability & Disaster Recovery Strategy")
    doc.add_paragraph(
        "EMS Pro is architected to scale horizontally. Stateless Uvicorn worker nodes can be added behind load balancers without session affinity. "
        "Database read-replicas support read-heavy dashboard and report queries. "
        "Point-in-time recovery (PITR) with continuous WAL archiving enables a Recovery Point Objective (RPO) < 5 minutes "
        "and Recovery Time Objective (RTO) < 30 minutes in case of catastrophic infrastructure failure."
    )

    output_path = "/Users/apple/Downloads/EMS/document/2_Architecture_Documentation.docx"
    doc.save(output_path)
    print(f"✔ Successfully generated: {output_path}")

if __name__ == "__main__":
    build_architecture_documentation()
