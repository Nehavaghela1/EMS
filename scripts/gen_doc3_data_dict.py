"""
Generate Document 3: Data Dictionary for EMS Pro.
Introspects all 38 SQLAlchemy models and writes comprehensive schema reference tables.
Outputs to: /Users/apple/Downloads/EMS/document/3_Data_Dictionary.docx
"""
import os
import sys
from docx.shared import Inches, Pt, RGBColor
from scripts.docx_helper import (
    create_styled_document, add_cover_page, add_header_1, add_header_2,
    add_header_3, add_callout, add_styled_table, COLOR_DARK_TEXT
)

# Import all models to populate Base.metadata
from app.db.base import Base
import app.modules.identity.models
import app.modules.hr.models
import app.modules.time_leave.models
import app.modules.payroll.models
import app.modules.performance.models
import app.modules.projects.models
import app.modules.platform.models

TABLE_DOMAIN_MAPPING = {
    "Identity & Access Management": [
        "companies", "company_settings", "users", "refresh_tokens", "industry_presets"
    ],
    "Core HR & Organization": [
        "departments", "employees"
    ],
    "Time, Attendance & Leave Management": [
        "shifts", "employee_shifts", "holidays", "attendance", "leave_types", "leave_balances", "leaves"
    ],
    "Statutory Payroll Engine": [
        "statutory_configs", "salary_structures", "salary_components", "employee_salaries",
        "payroll_runs", "payroll_items", "reimbursements", "pt_slabs", "tax_slabs"
    ],
    "Performance Management": [
        "performance_cycles", "performance_goals", "performance_reviews", "performance_summaries"
    ],
    "Projects & Workspace Management": [
        "projects", "project_members", "tasks", "task_comments", "time_entries", "milestones"
    ],
    "Platform Services & Audit": [
        "announcements", "employee_documents", "file_objects", "audit_logs", "notifications"
    ],
}

def get_column_type_str(col):
    t = str(col.type)
    if "VARCHAR" in t or "String" in t:
        return t.lower()
    if "UUID" in t:
        return "UUID"
    if "NUMERIC" in t:
        return f"Numeric({col.type.precision},{col.type.scale})"
    if "DATETIME" in t or "TIMESTAMP" in t:
        return "DateTime (UTC)"
    if "DATE" in t:
        return "Date"
    if "BOOLEAN" in t:
        return "Boolean"
    if "TEXT" in t:
        return "Text"
    if "INTEGER" in t:
        return "Integer"
    if "JSON" in t:
        return "JSON / JSONB"
    return t

def build_data_dictionary():
    doc = create_styled_document()
    add_cover_page(
        doc,
        title="EMS Pro — Data Dictionary",
        subtitle="Complete Relational Database Schema, Entity Attributes, Constraints & RLS Policies",
        doc_type="Database Schema & Data Dictionary Reference"
    )

    add_header_1(doc, "1. Schema Overview & Database Standards")
    doc.add_paragraph(
        "This Data Dictionary provides a comprehensive, field-by-field reference for all 38 relational tables "
        "comprising the EMS Pro database schema in PostgreSQL 16. "
        "Each table schema includes column data types, nullability, primary/foreign key relationships, "
        "unique constraints, default values, and operational descriptions."
    )

    add_callout(
        doc,
        "Every tenant table incorporates a foreign key 'company_id REFERENCES companies(id) ON DELETE CASCADE'. "
        "PostgreSQL kernel-level Row-Level Security (RLS) is enabled and forced on all tenant tables to prevent cross-tenant data leakage.",
        title="TENANT ISOLATION ARCHITECTURE"
    )

    overview_headers = ["Domain", "Table Name", "Total Columns", "Multi-Tenancy (RLS)", "Primary Business Role"]
    overview_data = []

    for domain_name, tables in TABLE_DOMAIN_MAPPING.items():
        for t_name in tables:
            if t_name in Base.metadata.tables:
                tbl = Base.metadata.tables[t_name]
                has_cid = "company_id" in tbl.columns
                rls_status = "Enabled (Tenant)" if has_cid and t_name != "users" else "System / Global"
                overview_data.append([domain_name.split()[0], t_name, str(len(tbl.columns)), rls_status, f"Manages {t_name.replace('_', ' ')}"])

    add_styled_table(doc, overview_headers, overview_data, [Inches(1.2), Inches(1.8), Inches(0.7), Inches(1.3), Inches(1.5)])

    # Detailed Table Specifications grouped by domain
    table_headers = ["Column Name", "Data Type", "Null?", "Key / Constraints", "Default / Ref", "Description"]
    col_widths = [Inches(1.5), Inches(1.1), Inches(0.5), Inches(1.1), Inches(0.9), Inches(1.4)]

    for domain_name, tables in TABLE_DOMAIN_MAPPING.items():
        add_header_1(doc, f"Domain: {domain_name}")

        for t_name in tables:
            if t_name not in Base.metadata.tables:
                continue

            tbl = Base.metadata.tables[t_name]
            add_header_2(doc, f"Table: {t_name}")

            doc_p = doc.add_paragraph()
            doc_p.add_run(f"Description: Primary relational entity for {t_name.replace('_', ' ')}. ")
            if "company_id" in tbl.columns and t_name != "users":
                doc_p.add_run("Protected by PostgreSQL Row-Level Security (RLS) policy.").font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            else:
                doc_p.add_run("Platform/Reference table managed with system-level constraints.")

            col_data = []
            for col in tbl.columns:
                c_name = col.name
                c_type = get_column_type_str(col)
                c_null = "NO" if not col.nullable else "YES"

                # Constraints
                constraints = []
                if col.primary_key:
                    constraints.append("PK")
                if col.unique:
                    constraints.append("UK")
                if col.foreign_keys:
                    for fk in col.foreign_keys:
                        constraints.append(f"FK -> {fk.target_fullname}")
                c_const = ", ".join(constraints) if constraints else "—"

                # Default
                c_def = "—"
                if col.default is not None:
                    c_def = str(col.default.arg) if hasattr(col.default, 'arg') else str(col.default)
                elif col.server_default is not None:
                    c_def = str(col.server_default.arg)

                # Meaningful description synthesis
                c_desc = f"{c_name.replace('_', ' ').capitalize()}"
                if c_name == "id":
                    c_desc = "Primary key UUID identifier"
                elif c_name == "company_id":
                    c_desc = "Tenant identifier for multi-tenant RLS isolation"
                elif c_name == "created_at":
                    c_desc = "Record creation UTC timestamp"
                elif c_name == "updated_at":
                    c_desc = "Last modification UTC timestamp"
                elif c_name == "status":
                    c_desc = "Operational lifecycle state machine value"
                elif "amount" in c_name or "salary" in c_name or "ctc" in c_name or "budget" in c_name:
                    c_desc = "Monetary currency amount in INR / tenant currency"

                col_data.append([c_name, c_type, c_null, c_const, c_def[:15], c_desc])

            add_styled_table(doc, table_headers, col_data, col_widths)

    output_path = "/Users/apple/Downloads/EMS/document/3_Data_Dictionary.docx"
    doc.save(output_path)
    print(f"✔ Successfully generated: {output_path}")

if __name__ == "__main__":
    build_data_dictionary()
